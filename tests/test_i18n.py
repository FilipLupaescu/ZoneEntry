import string

import pytest

import i18n
import zonelogic as zl
from i18n import LANGUAGES, STRINGS, tr

CODES = list(LANGUAGES)
# Strings that are legitimately identical in several languages
SAME_IS_FINE = {"col.n", "dlg.ok", "key.ctrl", "key.del", "dlg.import.title", "dlg.export.title"}


@pytest.fixture(autouse=True)
def _reset_language():
    yield
    i18n.set_language("en")


def placeholders(text: str) -> set[str]:
    return {name for _, name, _, _ in string.Formatter().parse(text) if name}


def test_five_languages():
    assert CODES == ["en", "ro", "hu", "pl", "de"]
    assert LANGUAGES["ro"] == "Română" and LANGUAGES["de"] == "Deutsch"


@pytest.mark.parametrize("key", sorted(STRINGS))
def test_every_key_has_every_language_with_same_placeholders(key):
    entry = STRINGS[key]
    assert set(entry) == set(CODES), f"{key}: languages {sorted(entry)}"
    expected = placeholders(entry["en"])
    for code in CODES:
        text = entry[code]
        assert text.strip(), f"{key}[{code}] is empty"
        assert placeholders(text) == expected, f"{key}[{code}] placeholders differ: {placeholders(text)} vs {expected}"
        text.format(**{p: "x" for p in expected})  # must be formattable


@pytest.mark.parametrize("key", sorted(set(STRINGS) - SAME_IS_FINE))
def test_translations_are_not_copies_of_english(key):
    entry = STRINGS[key]
    same = [c for c in CODES if c != "en" and entry[c] == entry["en"]]
    # a couple of short technical words can match in one language; never in all of them
    assert len(same) < len(CODES) - 1, f"{key} is untranslated"


def test_no_stray_english_in_other_languages_for_buttons():
    for key in ("btn.add", "btn.save", "btn.edit", "btn.delete", "btn.undo", "btn.new", "box.title"):
        for code in CODES[1:]:
            assert STRINGS[key][code] != STRINGS[key]["en"], (key, code)


def test_tr_switches_and_formats():
    i18n.set_language("de")
    assert tr("btn.add") == "Hinzufügen"
    assert tr("hint.editing", n=3) == "Zeile 3 wird bearbeitet – Enter speichert, Esc bricht ab."
    i18n.set_language("ro")
    assert tr("btn.delete") == "Șterge"
    i18n.set_language("hu")
    assert tr("lbl.end") == "Zárózóna"
    i18n.set_language("pl")
    assert tr("btn.export") == "Eksportuj do Excela…"


def test_unknown_language_and_key_fall_back():
    assert i18n.set_language("xx") == "en"
    assert tr("btn.add") == "Add"
    assert tr("'from'") == "'from'"  # plain text passes through (used for Excel header names)


def test_zone_errors_are_translated_and_follow_language_switch():
    with pytest.raises(zl.ZoneError) as ei:
        zl.normalize_zone("abc", "lbl.start")
    err = ei.value
    i18n.set_language("en")
    assert str(err) == "Start Zone must contain digits only (got 'abc')."
    i18n.set_language("de")  # same error object, new language
    assert str(err) == "Startzone: nur Ziffern erlaubt („abc“)."
    i18n.set_language("pl")
    assert "Strefa początkowa" in str(err)


def test_import_row_problem_is_localized(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["from", "to", "Division"])
    ws.append(["0100", "0050", "Kassen"])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    i18n.set_language("ro")
    _, problems = zl.import_xlsx(p)
    assert problems == ["Rândul 2: „from” este mai mare decât „to”."]


def test_detect_language_from_environment(monkeypatch):
    for var in ("LC_ALL", "LC_MESSAGES", "LANG"):
        monkeypatch.delenv(var, raising=False)
    assert i18n._language_from_environment() == "en"
    for env, code in (("ro_RO.UTF-8", "ro"), ("hu_HU", "hu"), ("pl_PL.UTF-8", "pl"), ("de_DE.UTF-8", "de"), ("fr_FR", "en")):
        monkeypatch.setenv("LANG", env)
        assert i18n._language_from_environment() == code


@pytest.mark.parametrize(
    "langid, code",
    [(0x0407, "de"), (0x0C07, "de"), (0x0807, "de"), (0x0418, "ro"), (0x040E, "hu"), (0x0415, "pl"), (0x0409, "en"), (0x040C, "en")],
)
def test_windows_ui_language_ids(langid, code):
    # de-DE, de-AT, de-CH, ro-RO, hu-HU, pl-PL, en-US, fr-FR
    assert i18n._language_from_windows_ui(langid) == code


def test_detect_language_never_crashes():
    assert i18n.detect_language() in LANGUAGES


def test_settings_roundtrip(tmp_path):
    assert zl.load_settings(tmp_path) == {"language": ""}
    zl.save_settings(tmp_path, {"language": "hu"})
    assert zl.load_settings(tmp_path)["language"] == "hu"
    (tmp_path / "settings.json").write_text('{"language": "klingon"}', encoding="utf-8")
    assert zl.load_settings(tmp_path)["language"] == ""
    (tmp_path / "settings.json").write_text("garbage", encoding="utf-8")
    assert zl.load_settings(tmp_path)["language"] == ""
