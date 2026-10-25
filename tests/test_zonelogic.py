from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

import zonelogic as zl
from zonelogic import Entry, ZoneError

SAMPLE = Path(__file__).parent / "data" / "ZoneUpdate_sample.xlsx"


# ---------------------------------------------------------------- zone numbers
@pytest.mark.parametrize(
    "text, expected",
    [("1", "0001"), ("42", "0042"), ("0042", "0042"), (" 9999 ", "9999"), ("007", "0007")],
)
def test_normalize_ok(text, expected):
    assert zl.normalize_zone(text) == expected


@pytest.mark.parametrize("text", ["", "   ", "abc", "12a", "-5", "0", "0000", "10000", "٣"])
def test_normalize_rejects(text):
    with pytest.raises(ZoneError):
        zl.normalize_zone(text)


@pytest.mark.parametrize(
    "end, expected",
    [("0042", "0043"), ("0001", "0002"), ("0999", "1000"), ("6393", "6394"), ("9998", "9999"), ("9999", None)],
)
def test_next_start_is_end_plus_one_with_leading_zeros(end, expected):
    assert zl.next_start(end) == expected


def test_suggest_start_uses_last_row():
    assert zl.suggest_start([]) is None
    rows = [Entry("0001", "0010", "Kassen"), Entry("0011", "0042", "Holz")]
    assert zl.suggest_start(rows) == "0043"


# ------------------------------------------------------------- overlaps / gaps
def test_overlap_detection():
    rows = [Entry("0001", "0010", "A"), Entry("0011", "0020", "B")]
    assert zl.find_overlap(rows, "0021", "0030") is None  # touching is fine
    assert zl.find_overlap(rows, "0010", "0012") == 0
    assert zl.find_overlap(rows, "0005", "0015") == 0
    assert zl.find_overlap(rows, "0001", "0030") == 0
    assert zl.find_overlap(rows, "0011", "0020", skip_index=1) is None  # editing row 1 itself


def test_gaps():
    rows = [Entry("0001", "0010", "A"), Entry("0013", "0020", "B"), Entry("0021", "0021", "C")]
    assert zl.find_gaps(rows) == [(11, 12)]
    assert zl.find_gaps([]) == []
    # an overlapping long range must not hide a real gap later on
    rows = [Entry("0001", "0100", "A"), Entry("0010", "0020", "B"), Entry("0150", "0160", "C")]
    assert zl.find_gaps(rows) == [(101, 149)]


def test_series_jump_into_new_thousand_block_is_not_a_gap():
    rows = [Entry("6001", "6397", "A"), Entry("7001", "7010", "B")]  # 6xxx -> 7xxx like your sheet
    assert zl.find_gaps(rows) == []
    # ...but a small skip right before a x001 block is still reported
    assert zl.find_gaps([Entry("6001", "6397", "A"), Entry("6401", "6410", "B")]) == [(6398, 6400)]
    assert zl.find_gaps([Entry("6001", "6990", "A"), Entry("7001", "7010", "B")]) == [(6991, 7000)]
    assert zl.skipped_between(6397, 7001) is None
    assert zl.skipped_between(7747, 7750) == (7748, 7749)
    assert zl.skipped_between(10, 11) is None


def test_sample_sheet_only_has_the_real_gap():
    entries, _ = zl.import_xlsx(SAMPLE)
    assert zl.find_gaps(entries) == [(7748, 7749)]


# ----------------------------------------------------------------- store zones
ZONES = list(zl.DEFAULT_STORE_ZONES)


def test_default_zones_are_unique_and_cover_the_sample():
    assert len({z.casefold() for z in ZONES}) == len(ZONES)
    wb = load_workbook(SAMPLE)
    used = {row[2] for row in wb.active.iter_rows(min_row=2, max_col=3, values_only=True)}
    assert used <= set(ZONES)


def test_search_prefix_first_letters():
    assert zl.filter_zones("san", ZONES)[0] == "Sanitär"
    assert zl.filter_zones("ELEK", ZONES) == ["Elektro"]
    assert zl.filter_zones("ka", ZONES)[:3] == ["Kassen", "Kalthalle Wand", "Kalthalle"]


def test_search_exact_match_beats_longer_names():
    assert zl.filter_zones("Garten", ZONES)[0] == "Garten"
    assert zl.filter_zones("Holz", ZONES)[0] == "Holz"
    assert zl.filter_zones("Kalthalle", ZONES)[0] == "Kalthalle"
    assert zl.filter_zones("Außenbereich", ZONES)[0] == "Außenbereich"


def test_search_german_spellings():
    assert zl.filter_zones("sanitaer", ZONES)[0] == "Sanitär"
    assert zl.filter_zones("sanitar", ZONES)[0] == "Sanitär"
    assert zl.filter_zones("aussen", ZONES)[0] == "Außenbereich Wand"  # first in list order
    assert zl.filter_zones("ausse", ZONES)[:2] == ["Außenbereich Wand", "Außenbereich"]


def test_search_word_start_and_substring():
    assert zl.filter_zones("fli", ZONES) == ["Laminat\\Fliesen"]
    assert set(zl.filter_zones("wand", ZONES)) == {"Garten Wand", "Kalthalle Wand", "Außenbereich Wand"}
    assert "Windfang" not in zl.filter_zones("wand", ZONES)
    assert zl.filter_zones("zzz", ZONES) == []
    assert zl.filter_zones("", ZONES) == ZONES


def test_resolve_zone():
    assert zl.resolve_zone("holz", ZONES) == "Holz"
    assert zl.resolve_zone("  AUSSENBEREICH ", ZONES) == "Außenbereich"  # ss == ß
    assert zl.resolve_zone("außenbereich", ZONES) == "Außenbereich"
    assert zl.resolve_zone("Hol", ZONES) is None
    assert zl.resolve_zone("", ZONES) is None


# ---------------------------------------------------------------------- Excel
def test_export_matches_original_layout(tmp_path):
    entries, problems = zl.import_xlsx(SAMPLE)
    assert problems == []
    assert len(entries) == 32

    out = tmp_path / "out.xlsx"
    zl.export_xlsx(out, entries)

    a, b = load_workbook(SAMPLE).active, load_workbook(out).active
    assert b.title == a.title == "Tabelle1"
    assert b.max_row == a.max_row == 33

    for r in range(1, a.max_row + 1):
        for c in (1, 2, 3):
            ca, cb = a.cell(r, c), b.cell(r, c)
            assert cb.value == ca.value, (r, c)
            assert isinstance(cb.value, str)  # text, so leading zeros survive
            assert cb.number_format == ca.number_format == "@"
            assert (cb.font.name, cb.font.sz) == (ca.font.name, ca.font.sz) == ("Aptos Narrow", 11)

    for col in "ABC":
        assert b.column_dimensions[col].width == pytest.approx(a.column_dimensions[col].width)
        assert b.column_dimensions[col].number_format == "@"  # blank cells below are text too


def test_export_keeps_leading_zeros_and_umlauts(tmp_path):
    out = tmp_path / "z.xlsx"
    zl.export_xlsx(out, [Entry("0001", "0042", "Sanitär"), Entry("0043", "0043", "Außenbereich")])
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == ["from", "to", "Division"]
    assert [c.value for c in ws[2]] == ["0001", "0042", "Sanitär"]
    assert [c.value for c in ws[3]] == ["0043", "0043", "Außenbereich"]


def test_import_rejects_foreign_sheet(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Name", "Price"])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    with pytest.raises(ZoneError):
        zl.import_xlsx(p)


def test_import_reports_bad_rows_and_accepts_numbers(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["from", "to", "Division"])
    ws.append([43, 56.0, "Holz"])  # Excel turned them into numbers
    ws.append(["0100", "0050", "Kassen"])  # from > to
    ws.append(["abc", "0060", "Deko"])
    ws.append([None, None, None])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    entries, problems = zl.import_xlsx(p)
    assert entries == [Entry("0043", "0056", "Holz")]
    assert len(problems) == 2 and problems[0].startswith("Row 3")


def test_default_export_name():
    d = date(2026, 10, 24)
    assert zl.default_export_name("", d) == "ZoneUpdate_2026-10-24.xlsx"
    name = zl.default_export_name("Bauhaus Köln / Nord: 3?", d)
    assert not any(ch in name for ch in '<>:"/\\|?*')
    assert name.endswith("_2026-10-24.xlsx") and "Bauhaus_Köln" in name


# ---------------------------------------------------------------- persistence
def test_zone_list_roundtrip_and_seeding(tmp_path):
    assert zl.load_zones(tmp_path) == ZONES  # first run seeds the file
    assert (tmp_path / "zones.txt").exists()
    zl.save_zones(tmp_path, ["Außenbereich", "Neue Zone", "neue zone", "# comment", ""])
    (tmp_path / "zones.txt").write_text("﻿Außenbereich\n\n# comment\nNeue Zone\nneue zone\n", encoding="utf-8")
    assert zl.load_zones(tmp_path) == ["Außenbereich", "Neue Zone"]


def test_session_roundtrip_and_corruption(tmp_path):
    rows = [Entry("0001", "0010", "Kassen"), Entry("0011", "0011", "Sanitär")]
    zl.save_session(tmp_path, "Markt Köln", rows, "/tmp")
    s = zl.load_session(tmp_path)
    assert s["store"] == "Markt Köln" and s["entries"] == rows and s["last_dir"] == "/tmp"

    (tmp_path / "session.json").write_text("{not json", encoding="utf-8")
    assert zl.load_session(tmp_path) == {"store": "", "last_dir": "", "entries": []}
