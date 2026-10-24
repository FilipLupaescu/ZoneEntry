"""Pure logic for ZoneEntry (no GUI imports, so it can be unit-tested anywhere).

Covers: zone-number validation/padding, the +1 "next start" rule, fuzzy
store-zone search, overlap/gap detection, saving the working session, and
reading/writing the Excel layout.
"""
from __future__ import annotations

import json
import os
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

APP_NAME = "ZoneEntry"

# --- Inventory zone rules ----------------------------------------------------
MIN_ZONE = 1
MAX_ZONE = 9999

# --- Excel layout (copied from your ZoneUpdate.xlsx) -------------------------
SHEET_NAME = "Tabelle1"
HEADERS = ("from", "to", "Division")
COLUMN_WIDTHS = (11.0, 11.85546875, 23.140625)  # A, B, C
FONT_NAME = "Aptos Narrow"
FONT_SIZE = 11
TEXT_FORMAT = "@"  # all three columns are formatted as Text, so 0042 stays 0042

# The Store Zones that appear in ZoneUpdate.xlsx (column "Division"), in order of
# first appearance. On first start they are written to zones.txt, which you can
# then edit from the app ("Zone list...").
DEFAULT_STORE_ZONES: tuple[str, ...] = (
    "Windfang",
    "Kassen",
    "Mittelgang",
    "Sanitär",
    "Elektro",
    "Leuchten",
    "Laminat\\Fliesen",
    "Haushalt",
    "Eisenwaren",
    "Holz",
    "Werkzeuge",
    "Auto",
    "Maschinen",
    "Bauelemente",
    "Baustoffe",
    "Deko",
    "Farben",
    "Weihnachtsmarkt",
    "Garten Grill",
    "Garten Wand",
    "Garten",
    "Kalthalle Wand",
    "Kalthalle",
    "Außenbereich Wand",
    "Außenbereich",
    "Eingang",
)


class ZoneError(ValueError):
    """A validation problem whose message is safe to show to the user as-is."""


@dataclass(frozen=True)
class Entry:
    """One row: a Store Zone and the range of Inventory Zones inside it."""

    start: str  # "0043"
    end: str  # "0056"
    zone: str  # "Mittelgang"

    @property
    def first(self) -> int:
        return int(self.start)

    @property
    def last(self) -> int:
        return int(self.end)

    @property
    def count(self) -> int:
        return self.last - self.first + 1

    def label(self) -> str:
        return f"{self.start}–{self.end}  {self.zone}"


# =============================================================================
# Zone numbers
# =============================================================================
def pad(n: int) -> str:
    """42 -> '0042'."""
    return f"{n:04d}"


def fmt_range(a: int, b: int) -> str:
    return pad(a) if a == b else f"{pad(a)}–{pad(b)}"


def normalize_zone(text: str, what: str = "Zone") -> str:
    """Validate user input and return it padded to 4 digits ('42' -> '0042')."""
    t = (text or "").strip()
    if not t:
        raise ZoneError(f"{what} is empty.")
    if not (t.isascii() and t.isdigit()):
        raise ZoneError(f"{what} must contain digits only (got '{t}').")
    if len(t) > 4:
        raise ZoneError(f"{what} can have at most 4 digits (got '{t}').")
    n = int(t)
    if not MIN_ZONE <= n <= MAX_ZONE:
        raise ZoneError(f"{what} must be between {pad(MIN_ZONE)} and {pad(MAX_ZONE)}.")
    return pad(n)


def next_start(end: str) -> str | None:
    """The automatic suggestion: previous End + 1 ('0042' -> '0043').

    Returns None when the previous End is already 9999.
    """
    n = int(end) + 1
    return pad(n) if n <= MAX_ZONE else None


def suggest_start(entries: Sequence[Entry]) -> str | None:
    """Suggested Start Zone for the next row, based on the last row's End."""
    if not entries:
        return None
    return next_start(entries[-1].end)


def find_overlap(
    entries: Sequence[Entry], start: str, end: str, skip_index: int | None = None
) -> int | None:
    """Index of the first existing row whose range overlaps start..end."""
    s, e = int(start), int(end)
    for i, other in enumerate(entries):
        if i == skip_index:
            continue
        if s <= other.last and other.first <= e:
            return i
    return None


SERIES_JUMP_MIN = 100  # a skip this big *into* an x001 block is a new series, not a mistake


def is_series_jump(gap_start: int, next_start: int) -> bool:
    """True for deliberate jumps such as 6397 -> 7001 (a new thousand-block begins)."""
    return next_start % 1000 == 1 and next_start - gap_start >= SERIES_JUMP_MIN


def skipped_between(prev_end: int, start: int) -> tuple[int, int] | None:
    """Zone numbers silently skipped between a row's End and the next row's Start."""
    if start > prev_end + 1 and not is_series_jump(prev_end + 1, start):
        return prev_end + 1, start - 1
    return None


def find_gaps(entries: Iterable[Entry]) -> list[tuple[int, int]]:
    """Inventory-zone numbers between the lowest and highest zone not covered by any row.

    Jumps into a new thousand-block (6397 -> 7001) are not reported.
    """
    gaps: list[tuple[int, int]] = []
    reach: int | None = None
    for a, b in sorted((e.first, e.last) for e in entries):
        if reach is not None:
            gap = skipped_between(reach, a)
            if gap:
                gaps.append(gap)
        reach = b if reach is None else max(reach, b)
    return gaps


# =============================================================================
# Store-zone search ("type the first few characters")
# =============================================================================
_UMLAUTS = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue"})
_WORD_SPLIT = re.compile(r"[^0-9a-z]+")


def _variants(text: str) -> set[str]:
    """Search keys: case-folded (ß->ss), umlauts as 'ae' or stripped to 'a'."""
    s = text.casefold()
    expanded = s.translate(_UMLAUTS)
    stripped = "".join(
        ch for ch in unicodedata.normalize("NFKD", s) if not unicodedata.combining(ch)
    )
    return {expanded, stripped}


def _score(query: str, name: str) -> int | None:
    if name == query:
        return -1  # exact match always wins ("Garten" before "Garten Grill")
    if name.startswith(query):
        return 0
    if any(w.startswith(query) for w in _WORD_SPLIT.split(name) if w):
        return 1  # "wand" finds "Garten Wand"
    if query in name:
        return 2
    return None


def filter_zones(query: str, zones: Iterable[str]) -> list[str]:
    """Zones matching `query`, best first. Empty query returns everything."""
    zones = list(zones)
    q = (query or "").strip()
    if not q:
        return zones
    q_keys = _variants(q)
    ranked: list[tuple[int, int, str]] = []
    for i, name in enumerate(zones):
        best: int | None = None
        for nk in _variants(name):
            for qk in q_keys:
                s = _score(qk, nk)
                if s is not None and (best is None or s < best):
                    best = s
        if best is not None:
            ranked.append((best, i, name))
    ranked.sort()
    return [name for _, _, name in ranked]


def resolve_zone(text: str, zones: Iterable[str]) -> str | None:
    """Exact (case-insensitive) match against the list, returning the canonical name."""
    t = (text or "").strip().casefold()
    if not t:
        return None
    for z in zones:
        if z.casefold() == t:
            return z
    return None


# =============================================================================
# Files on disk: zone list + autosaved session
# =============================================================================
def app_data_dir() -> Path:
    override = os.environ.get("ZONEENTRY_HOME")
    if override:
        p = Path(override)
    elif sys.platform == "win32":
        base = os.environ.get("APPDATA") or str(Path.home() / "AppData" / "Roaming")
        p = Path(base) / APP_NAME
    elif sys.platform == "darwin":
        p = Path.home() / "Library" / "Application Support" / APP_NAME
    else:
        base = os.environ.get("XDG_DATA_HOME") or str(Path.home() / ".local" / "share")
        p = Path(base) / APP_NAME
    p.mkdir(parents=True, exist_ok=True)
    return p


def parse_zone_lines(text: str) -> list[str]:
    """One zone per line; blank lines, '#' comments and duplicates are ignored."""
    seen: set[str] = set()
    out: list[str] = []
    for line in text.splitlines():
        name = line.strip()
        if not name or name.startswith("#"):
            continue
        key = name.casefold()
        if key not in seen:
            seen.add(key)
            out.append(name)
    return out


def save_zones(folder: Path, zones: Sequence[str]) -> None:
    (folder / "zones.txt").write_text("\n".join(zones) + "\n", encoding="utf-8")


def load_zones(folder: Path) -> list[str]:
    path = folder / "zones.txt"
    if not path.exists():
        save_zones(folder, DEFAULT_STORE_ZONES)
        return list(DEFAULT_STORE_ZONES)
    try:
        zones = parse_zone_lines(path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeDecodeError):
        zones = []
    return zones or list(DEFAULT_STORE_ZONES)


def save_session(
    folder: Path, store: str, entries: Sequence[Entry], last_dir: str = ""
) -> None:
    """Autosave so a crash or closed window never loses a count."""
    data = {
        "version": 1,
        "store": store,
        "last_dir": last_dir,
        "entries": [[e.start, e.end, e.zone] for e in entries],
    }
    path = folder / "session.json"
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, path)


def load_session(folder: Path) -> dict:
    """Returns {'store', 'last_dir', 'entries'}; falls back to empty on any problem."""
    result: dict = {"store": "", "last_dir": "", "entries": []}
    try:
        data = json.loads((folder / "session.json").read_text(encoding="utf-8"))
        result["store"] = str(data.get("store", ""))
        result["last_dir"] = str(data.get("last_dir", ""))
        for s, e, z in data.get("entries", []):
            s, e = normalize_zone(s), normalize_zone(e)
            if int(s) <= int(e):
                result["entries"].append(Entry(s, e, str(z)))
    except (OSError, ValueError, TypeError):  # ZoneError is a ValueError
        pass
    return result


# =============================================================================
# Excel
# =============================================================================
def default_export_name(store: str, today: date | None = None) -> str:
    safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", store or "").strip().strip(".")
    safe = re.sub(r"\s+", "_", safe)
    stamp = (today or date.today()).isoformat()
    return "_".join(["ZoneUpdate", *([safe] if safe else []), stamp]) + ".xlsx"


def export_xlsx(path: str | os.PathLike, entries: Sequence[Entry]) -> None:
    """Write the sheet exactly like ZoneUpdate.xlsx: from | to | Division, all text."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    font = Font(name=FONT_NAME, size=FONT_SIZE)

    for letter, width in zip("ABC", COLUMN_WIDTHS):
        col = ws.column_dimensions[letter]
        col.width = width
        col.number_format = TEXT_FORMAT  # column-level style, as in the original
        col.font = font

    rows = [HEADERS] + [(e.start, e.end, e.zone) for e in entries]
    for r, values in enumerate(rows, start=1):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c)
            cell.value = value  # str -> stored as text, keeps leading zeros
            cell.number_format = TEXT_FORMAT
            cell.font = font
    wb.save(path)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def import_xlsx(path: str | os.PathLike) -> tuple[list[Entry], list[str]]:
    """Read a sheet in the same layout. Returns (entries, list of skipped-row problems)."""
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb.active
        header = [_cell_text(ws.cell(row=1, column=c).value).casefold() for c in (1, 2, 3)]
        if header != [h.casefold() for h in HEADERS]:
            raise ZoneError(
                "This file doesn't look like a zone sheet: "
                "expected the headers 'from', 'to', 'Division' in A1:C1."
            )
        entries: list[Entry] = []
        problems: list[str] = []
        for r, (a, b, c) in enumerate(
            ws.iter_rows(min_row=2, max_col=3, values_only=True), start=2
        ):
            if all(v in (None, "") for v in (a, b, c)):
                continue
            try:
                s = normalize_zone(_cell_text(a), "'from'")
                e = normalize_zone(_cell_text(b), "'to'")
                if int(s) > int(e):
                    raise ZoneError("'from' is bigger than 'to'.")
            except ZoneError as exc:
                problems.append(f"Row {r}: {exc}")
                continue
            entries.append(Entry(s, e, _cell_text(c)))
        return entries, problems
    finally:
        wb.close()
